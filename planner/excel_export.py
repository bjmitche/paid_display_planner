from __future__ import annotations

import math
from copy import copy
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from .economics import estimate_inventory
from .models import Activation, Campaign, Inventory, Product

TEMPLATE_PATH = Path(__file__).parents[1] / "templates" / "Paid_Display_Planner_Template.xlsx"


def _copy_row_style(sheet, source_row: int, target_row: int, max_column: int) -> None:
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def _write_rows(
    sheet, start_row: int, rows: list[list[Any]], template_row: int, max_column: int
) -> None:
    for offset, values in enumerate(rows):
        row = start_row + offset
        if row != template_row:
            _copy_row_style(sheet, template_row, row, max_column)
        for column, value in enumerate(values, 1):
            sheet.cell(row, column).value = value


def _replace_table(sheet, old_name: str, new_name: str, ref: str) -> None:
    if old_name in sheet.tables:
        del sheet.tables[old_name]
    table = Table(displayName=new_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _set_table_ref(sheet, name: str, ref: str) -> None:
    table = sheet.tables[name]
    table.ref = ref
    if table.autoFilter is not None:
        table.autoFilter.ref = ref


def _clear_data(sheet, start_row: int, end_row: int, max_column: int) -> None:
    for row in range(start_row, end_row + 1):
        for column in range(1, max_column + 1):
            sheet.cell(row, column).value = None


def _histogram(values: list[float], bins: int = 20) -> list[tuple[float, float, int]]:
    if not values:
        return [(0.0, 1.0, 0)] * bins
    low, high = min(values), max(values)
    width = (high - low) / bins if high > low else 1.0
    counts = [0] * bins
    for value in values:
        counts[min(bins - 1, math.floor((value - low) / width))] += 1
    return [(low + width * i, low + width * (i + 1), counts[i]) for i in range(bins)]


def _quartile(values: list[float], fraction: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    index = min(len(ordered) - 1, max(0, int(len(ordered) * fraction) - 1))
    return ordered[index]


def _median(values: list[float]) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    middle = len(ordered) // 2
    return ordered[middle] if len(ordered) % 2 else (ordered[middle - 1] + ordered[middle]) / 2


def build_workbook(
    campaign: Campaign,
    selected: list[Activation],
    inventories: dict[str, Inventory],
    products: dict[str, Product],
    product_overrides: dict[str, Product],
    product_by_activation: dict[str, Product],
    history_by_inventory: dict[str, list[Activation]],
    conversion_by_activation: dict[str, tuple[float, float, float, float]],
    target_currency: str,
    fx_rates: dict[str, float],
    iterations: int,
    activation_simulations: list[tuple[Activation, list[dict[str, float]]]],
    campaign_rows: list[dict[str, float]],
) -> bytes:
    workbook = load_workbook(TEMPLATE_PATH)
    inputs = workbook["Model Inputs & Assumptions"]
    summary = workbook["Performance Summary"]
    activation_detail = workbook["Activation Detail"]
    simulation_detail = workbook["Simulation Detail"]

    inputs["C5"] = campaign.name
    inputs["C6"] = "Streamlit scenario"
    inputs["C7"] = date.today()
    inputs["C8"] = target_currency
    inputs["C9"] = iterations
    inputs["C10"] = 42
    inputs["C12"] = "Exported from Paid Display Planner. Values are a read-only scenario snapshot."

    selected_inventory_ids = list(dict.fromkeys(a.inventory_id for a in selected if a.inventory_id))
    relevant_inventories = [inventories[i] for i in selected_inventory_ids if i in inventories]
    inventory_rows: list[list[Any]] = []
    for inventory in relevant_inventories:
        history = history_by_inventory.get(inventory.id, [])
        estimate = estimate_inventory(inventory, history, inventory.cost)
        valid_impressions = [
            a.actual_impressions for a in history if a.actual_impressions is not None
        ]
        valid_views = [a.actual_video_views for a in history if a.actual_video_views is not None]
        valid_clicks = [a.actual_clicks for a in history if a.actual_clicks is not None]
        hist_total_impressions = sum(valid_impressions)
        hist_total_views = sum(valid_views)
        hist_total_clicks = sum(valid_clicks)
        inventory_rows.append(
            [
                inventory.id,
                inventory.name,
                estimate.method,
                len(history),
                sum(a.actual_impressions is not None for a in history),
                sum(a.actual_video_views is not None for a in history),
                sum(a.actual_clicks is not None for a in history),
                _median(valid_impressions),
                hist_total_impressions,
                hist_total_views,
                hist_total_clicks,
                hist_total_views / hist_total_impressions if hist_total_impressions else 0,
                hist_total_clicks / hist_total_impressions if hist_total_impressions else 0,
                inventory.expected_impressions or 0,
                inventory.expected_view_rate or 0,
                inventory.expected_ctr or 0,
                inventory.impressions_sigma or 0,
                inventory.view_rate_sigma or 0,
                inventory.ctr_sigma or 0,
                1 if valid_impressions else 0,
                1 if valid_views else 0,
                1 if valid_clicks else 0,
                estimate.impressions,
                estimate.view_rate,
                estimate.ctr,
                estimate.method,
                "Historical target-rate estimate with Inventory fallback where required.",
            ]
        )

    used_product_ids = list(dict.fromkeys(product.id for product in product_by_activation.values()))
    used_products = [
        product_overrides.get(i) or products[i]
        for i in used_product_ids
        if i in products or i in product_overrides
    ]
    product_rows = [
        [
            p.id,
            p.name,
            p.currency,
            p.average_purchase_amount or 0,
            p.gross_margin_pct or 0,
            p.holding_period or 0,
            p.ltv or 0,
            fx_rates.get(p.currency or target_currency, 1.0),
            "Loaded from Notion with Streamlit override applied.",
        ]
        for p in used_products
    ]

    activation_rows: list[list[Any]] = []
    for activation in selected:
        inventory = inventories.get(activation.inventory_id or "")
        product = (
            product_by_activation.get(activation.id)
            or product_overrides.get(activation.product_id or "")
            or products.get(activation.product_id or "")
        )
        estimate = (
            estimate_inventory(
                inventory, history_by_inventory.get(inventory.id, []), activation.cost
            )
            if inventory
            else None
        )
        imp, view, click, sigma = conversion_by_activation[activation.id]
        activation_rows.append(
            [
                activation.id,
                activation.name,
                "Scenario activation"
                if activation.id.startswith("scenario:")
                else "Campaign activation",
                inventory.id if inventory else None,
                inventory.name if inventory else None,
                inventory.format if inventory else None,
                activation.status,
                inventory.rate_basis if inventory else None,
                activation.cost,
                activation.currency,
                fx_rates.get(activation.currency or target_currency, 1.0),
                product.id if product else None,
                product.name if product else None,
                imp,
                view,
                click,
                sigma,
                estimate.method if estimate else "Unavailable",
                estimate.impressions if estimate else None,
                estimate.impressions_sigma if estimate else None,
                estimate.view_rate if estimate else None,
                estimate.view_rate_sigma if estimate else None,
                estimate.ctr if estimate else None,
                estimate.ctr_sigma if estimate else None,
                "Exported scenario input.",
            ]
        )

    fx_rows = [
        [target_currency, target_currency, 1.0, "Identity", date.today(), "No", "Target currency"]
    ]
    fx_rows += [
        [
            c,
            target_currency,
            r,
            "Frankfurter/ECB preload",
            date.today(),
            "No",
            "Editable source rate",
        ]
        for c, r in fx_rates.items()
    ]
    extra_inventory = max(0, len(inventory_rows) - 7)
    extra_activations = max(0, len(activation_rows) - 8)
    extra_products = max(0, len(product_rows) - 3)
    extra_fx = max(0, len(fx_rows) - 4)
    # Insert from the bottom upward so earlier insertion does not invalidate anchors.
    if extra_fx:
        inputs.insert_rows(51, extra_fx)
    if extra_products:
        inputs.insert_rows(43, extra_products)
    if extra_activations:
        inputs.insert_rows(36, extra_activations)
    if extra_inventory:
        inputs.insert_rows(24, extra_inventory)
    activation_start = 28 + extra_inventory
    product_start = 40 + extra_inventory + extra_activations
    fx_start = 47 + extra_inventory + extra_activations + extra_products
    _clear_data(inputs, 17, 23 + extra_inventory, 27)
    _clear_data(inputs, activation_start, activation_start + 7 + extra_activations, 25)
    _clear_data(inputs, product_start, product_start + 2 + extra_products, 9)
    _clear_data(inputs, fx_start, fx_start + 3 + extra_fx, 7)
    _write_rows(inputs, 17, inventory_rows, 17, 27)
    _write_rows(inputs, activation_start, activation_rows, activation_start, 25)
    _write_rows(inputs, product_start, product_rows, product_start, 9)
    _write_rows(inputs, fx_start, fx_rows, fx_start, 7)
    _set_table_ref(
        inputs,
        "InventoryEstimates",
        f"A16:AA{16 + len(inventory_rows)}",
    )
    _set_table_ref(
        inputs,
        "Activations",
        f"A{activation_start - 1}:Y{activation_start - 1 + len(activation_rows)}",
    )
    _set_table_ref(
        inputs,
        "Products",
        f"A{product_start - 1}:I{product_start - 1 + len(product_rows)}",
    )
    _set_table_ref(
        inputs,
        "FXRates",
        f"A{fx_start - 1}:G{fx_start - 1 + len(fx_rows)}",
    )

    detail_headers = [
        "Iteration",
        "Activation",
        "Impressions",
        "View rate",
        "Views",
        "CTR",
        "Clicks",
        "Conversions",
        "Cost",
        "Flows",
        "LTV value",
        "ROI",
        "Cost per conversion",
        "Product",
        "Notes",
        "z — impressions",
        "z — view rate",
        "z — CTR",
        "z — conversion",
    ]
    detail_rows: list[list[Any]] = []
    for activation, rows in activation_simulations:
        product = (
            product_by_activation.get(activation.id)
            or product_overrides.get(activation.product_id or "")
            or products.get(activation.product_id or "")
        )
        for iteration, row in enumerate(rows, 1):
            detail_rows.append(
                [
                    iteration,
                    activation.name,
                    row["impressions"],
                    row["views"] / row["impressions"] if row["impressions"] else 0,
                    row["views"],
                    row["clicks"] / row["impressions"] if row["impressions"] else 0,
                    row["clicks"],
                    row["conversions"],
                    row["cost"],
                    row["flows"],
                    row["value"],
                    row["roi"],
                    row["cost"] / row["conversions"] if row["conversions"] else 0,
                    product.name if product else None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ]
            )
    _clear_data(activation_detail, 3, activation_detail.max_row, 19)
    _write_rows(activation_detail, 3, [detail_headers] + detail_rows, 3, 19)
    _replace_table(
        activation_detail, "ActivationDetail", "ActivationDetail", f"A3:S{3 + len(detail_rows)}"
    )

    product_ids = used_product_ids[:3]
    sim_headers = [
        "Iteration",
        "Total impressions",
        "Total views",
        "Total clicks",
        "Total conversions",
        "Total cost",
        "Total flows",
        "Total LTV value",
        "ROI",
        "Cost per conversion",
        "View rate %",
        "CTR %",
        "Notes",
    ]
    for pid in product_ids:
        sim_headers.extend(
            [
                f"{pid} conversions",
                f"{pid} flows",
                f"{pid} LTV value",
                f"{pid} cost",
                f"{pid} ROI",
                f"{pid} cost per conversion",
            ]
        )
    sim_headers += [f"Reserved {i}" for i in range(31 - len(sim_headers))]
    sim_rows: list[list[Any]] = []
    for i, campaign_row in enumerate(campaign_rows, 1):
        total_impressions = campaign_row["impressions"]
        total_views = campaign_row["views"]
        total_clicks = campaign_row["clicks"]
        values: list[Any] = [
            i,
            total_impressions,
            total_views,
            total_clicks,
            campaign_row["conversions"],
            campaign_row["cost"],
            campaign_row["flows"],
            campaign_row["value"],
            campaign_row["roi"],
            campaign_row["cost_per_conversion"],
            total_views / total_impressions if total_impressions else 0,
            total_clicks / total_impressions if total_impressions else 0,
            None,
        ]
        for pid in product_ids:
            pr = {"conversions": 0.0, "flows": 0.0, "value": 0.0, "cost": 0.0}
            for activation, rows in activation_simulations:
                if activation.product_id == pid:
                    row = rows[i - 1]
                    for key in pr:
                        pr[key] += row["value" if key == "value" else key]
            values += [
                pr["conversions"],
                pr["flows"],
                pr["value"],
                pr["cost"],
                pr["value"] / pr["cost"] if pr["cost"] else 0,
                pr["cost"] / pr["conversions"] if pr["conversions"] else 0,
            ]
        sim_rows.append(values + [None] * (31 - len(values)))
    _clear_data(simulation_detail, 3, simulation_detail.max_row, 31)
    _write_rows(simulation_detail, 3, [sim_headers] + sim_rows, 3, 31)
    _replace_table(
        simulation_detail, "SimulationDetail", "SimulationDetail", f"A3:AE{3 + len(sim_rows)}"
    )

    summary["B5"] = campaign.name
    summary["B6"] = "Streamlit scenario"
    summary["B7"] = target_currency
    summary["B8"] = len(selected)
    summary["B9"] = len(used_products)
    summary["B10"] = iterations
    metric_map = {
        "Impressions": "impressions",
        "Video views": "views",
        "Clicks": "clicks",
        "Conversions": "conversions",
        "Flows": "flows",
        "Cost": "cost",
        "Cost per conversion": "cost_per_conversion",
        "LTV value": "value",
        "ROI": "roi",
    }
    direction = {"cost": "lower", "cost_per_conversion": "lower"}
    for row, (label, key) in enumerate(metric_map.items(), 15):
        values = [r[key] for r in campaign_rows] or [0]
        q1, med, q3 = _quartile(values, 0.25), _quartile(values, 0.5), _quartile(values, 0.75)
        summary.cell(row, 1).value = label
        summary.cell(row, 2).value = q3 if direction.get(key) == "lower" else q1
        summary.cell(row, 3).value = med
        summary.cell(row, 4).value = q1 if direction.get(key) == "lower" else q3
        summary.cell(row, 5).value = (
            "Lower is better" if direction.get(key) == "lower" else "Higher is better"
        )

    activation_performance_headers = [
        "Activation name",
        "Source",
        "Inventory",
        "Product",
        "Estimation method",
        "Expected impressions",
        "Expected view rate %",
        "Expected CTR %",
        "Conversions — median",
        "Conversions — Q1",
        "Conversions — Q3",
        "Flows — median",
        "LTV value — median",
        "Cost — median",
        "ROI — median",
        "Cost per conversion — median",
    ]
    activation_performance_rows = []
    for activation, rows in activation_simulations:
        inventory = inventories.get(activation.inventory_id or "")
        product = product_by_activation.get(activation.id) or products.get(
            activation.product_id or ""
        )
        estimate = (
            estimate_inventory(
                inventory,
                history_by_inventory.get(inventory.id, []),
                activation.cost,
            )
            if inventory
            else None
        )
        activation_performance_rows.append(
            [
                activation.name,
                "Scenario activation"
                if activation.id.startswith("scenario:")
                else "Campaign activation",
                inventory.name if inventory else None,
                product.name if product else None,
                estimate.method if estimate else "Unavailable",
                estimate.impressions if estimate else None,
                estimate.view_rate if estimate else None,
                estimate.ctr if estimate else None,
                _quartile([row["conversions"] for row in rows], 0.50),
                _quartile([row["conversions"] for row in rows], 0.25),
                _quartile([row["conversions"] for row in rows], 0.75),
                _quartile([row["flows"] for row in rows], 0.50),
                _quartile([row["value"] for row in rows], 0.50),
                _quartile([row["cost"] for row in rows], 0.50),
                _quartile([row["roi"] for row in rows], 0.50),
                _quartile(
                    [row["cost"] / row["conversions"] if row["conversions"] else 0 for row in rows],
                    0.50,
                ),
            ]
        )
    extra_activation_performance = max(0, len(activation_performance_rows) - 8)
    if extra_activation_performance:
        summary.insert_rows(44, extra_activation_performance)
    _clear_data(
        summary,
        35,
        43 + extra_activation_performance,
        len(activation_performance_headers),
    )
    _write_rows(
        summary,
        35,
        [activation_performance_headers] + activation_performance_rows,
        35,
        len(activation_performance_headers),
    )
    _replace_table(
        summary,
        "ActivationPerformance",
        "ActivationPerformance",
        f"A35:P{35 + len(activation_performance_rows)}",
    )

    conversion_bins = _histogram([r["conversions"] for r in campaign_rows])
    roi_bins = _histogram([r["roi"] for r in campaign_rows])
    histogram_note_row = 83 + extra_activation_performance
    histogram_start_row = 85 + extra_activation_performance
    summary.cell(histogram_note_row, 1).value = (
        "Fixed histogram bins: 20 equal-width bins per distribution; "
        "final bin includes the maximum."
    )
    for offset, (start, end, count) in enumerate(conversion_bins, histogram_start_row):
        (
            summary.cell(offset, 1).value,
            summary.cell(offset, 2).value,
            summary.cell(offset, 3).value,
            summary.cell(offset, 4).value,
        ) = start, end, f"{start:,.2f}", count
        (
            summary.cell(offset, 6).value,
            summary.cell(offset, 7).value,
            summary.cell(offset, 8).value,
            summary.cell(offset, 9).value,
        ) = (
            roi_bins[offset - histogram_start_row][0],
            roi_bins[offset - histogram_start_row][1],
            f"{roi_bins[offset - histogram_start_row][0]:,.2f}",
            roi_bins[offset - histogram_start_row][2],
        )

    output = BytesIO()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output)
    return output.getvalue()
