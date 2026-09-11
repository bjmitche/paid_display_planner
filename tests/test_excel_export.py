from io import BytesIO

from openpyxl import load_workbook

from planner.excel_export import build_workbook
from planner.models import Activation, Campaign, Inventory, Product


def test_build_workbook_populates_template():
    activation = Activation("a", "Activation", "Done", "i", "p", 100, "EUR", None, None, None)
    inventory = Inventory(
        id="i",
        name="Inventory",
        format="Display",
        buying_model="Fixed placement fee",
        objective="Impressions",
        currency="EUR",
        cost=100,
        rate_basis="CPM",
        expected_cpm=None,
        expected_cpc=None,
        expected_cpv=None,
        cpm_sigma=None,
        cpc_sigma=None,
        cpv_sigma=None,
        expected_impressions=1000,
        expected_view_rate=0,
        expected_ctr=0.01,
        impressions_sigma=0.1,
        view_rate_sigma=0.1,
        ctr_sigma=0.1,
    )
    product = Product("p", "Product", 1000, None, "EUR", 0.01, 2)
    rows = [
        {
            "impressions": 1000.0,
            "views": 0.0,
            "clicks": 10.0,
            "conversions": 1.0,
            "cost": 100.0,
            "flows": 1000.0,
            "value": 20.0,
            "roi": 0.2,
        }
        for _ in range(3)
    ]
    campaign_rows = [{**row, "cost_per_conversion": 100.0} for row in rows]
    workbook = build_workbook(
        Campaign("c", "Campaign", ("a",), None, None, "EUR"),
        [activation],
        {"i": inventory},
        {"p": product},
        {"p": product},
        {"i": [activation]},
        {"a": (0.0, 0.0, 0.01, 0.1)},
        "EUR",
        {},
        3,
        [(activation, rows)],
        campaign_rows,
    )
    assert workbook[:2] == b"PK"
    assert len(workbook) > 100_000
    loaded = load_workbook(BytesIO(workbook), data_only=False)
    assert loaded.sheetnames == [
        "Model Inputs & Assumptions",
        "Performance Summary",
        "Activation Detail",
        "Simulation Detail",
    ]
    assert loaded["Model Inputs & Assumptions"]["C5"].value == "Campaign"
    assert loaded["Model Inputs & Assumptions"].tables["Activations"].ref == "A27:Y28"
    assert loaded["Activation Detail"].tables["ActivationDetail"].ref == "A3:S6"
    assert loaded["Simulation Detail"].tables["SimulationDetail"].ref == "A3:AE6"
    assert sum(loaded["Performance Summary"].cell(row, 4).value for row in range(85, 105)) == 3
    assert sum(loaded["Performance Summary"].cell(row, 9).value for row in range(85, 105)) == 3
    assert (
        loaded["Performance Summary"]["B24"].value
        == "=QUARTILE('Simulation Detail'!$H$4:$H$503,1)"
    )
    assert (
        loaded["Performance Summary"]["C25"].value
        == "=QUARTILE('Simulation Detail'!$I$4:$I$503,2)"
    )
